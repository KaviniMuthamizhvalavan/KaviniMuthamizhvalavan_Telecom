"""Build presentation/Project_Presentation.pptx (plan §12.7: 8-minute talk with live demo, then Q&A).

Run from the repo root:  python -m scripts.build_presentation   (needs python-pptx; build-time only)
Every figure on the slides comes from evaluation/results/final/ (Sonnet 5.5 run) or
evaluation/results/runs/aicredits_haiku_demo/ (demo-setup run). Updated 2026-10-11 for the React frontend and both runs.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from src.utils.config import PROJECT_ROOT

OUT = PROJECT_ROOT / "presentation" / "Project_Presentation.pptx"
NAVY, ORANGE, GREEN, GREY, INK = (RGBColor(0x0B, 0x3D, 0x5C), RGBColor(0xC0, 0x61, 0x2B), RGBColor(0x2E, 0x7D, 0x4F),
                                  RGBColor(0x6B, 0x78, 0x83), RGBColor(0x1D, 0x23, 0x29))
FONT = "Segoe UI"
W, H = Inches(13.333), Inches(7.5)


def text(frame, lines, size=18, color=INK, bold_first=False, space=6):
    frame.word_wrap = True
    frame.clear()
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        level = 1 if line.startswith("  ") else 0
        p.level = level
        run = p.add_run()
        run.text = ("• " if level == 0 and not line.startswith("#") else "– ") + line.strip().lstrip("#").strip() \
            if not line.startswith("!") else line[1:]
        run.font.size, run.font.name, run.font.color.rgb = Pt(size - 2 * level), FONT, color
        run.font.bold = bold_first and i == 0
        p.space_after = Pt(space)


def slide(prs, title, notes="", subtitle=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = s.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.75)).text_frame
    run = tf.paragraphs[0].add_run()
    run.text, run.font.size, run.font.bold, run.font.name = title, Pt(30), True, FONT
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if subtitle:
        st = s.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.5)).text_frame
        r = st.paragraphs[0].add_run()
        r.text, r.font.size, r.font.name, r.font.color.rgb, r.font.italic = subtitle, Pt(15), FONT, GREY, True
    foot = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35)).text_frame
    r = foot.paragraphs[0].add_run()
    r.text, r.font.size, r.font.name, r.font.color.rgb = "AI-Powered Telecom Network Operations Assistant", Pt(10), FONT, GREY
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def box(s, x, y, w, h, title, body, fill, line, title_size=18, body_size=15, centered=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if centered else MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.15)
    align = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    tf.paragraphs[0].alignment = align
    r = tf.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.name, r.font.color.rgb = title, Pt(title_size), True, FONT, INK
    tf.paragraphs[0].space_after = Pt(8)
    for line in body:
        p = tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        rr = p.add_run()
        rr.text, rr.font.size, rr.font.name, rr.font.color.rgb = line, Pt(body_size), FONT, INK
    return shp


def arrow(s, x1, y1, x2, y2, color=INK, dashed=False):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    if dashed:
        c.line.dash_style = 7  # dash
    end = c.line._get_or_add_ln()
    end.append(end.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}))
    return c


def table(s, x, y, w, rows, col_w=None, size=13, h_row=0.42):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h_row * len(rows)))
    t = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), FONT
                    run.font.bold = r == 0
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (RGBColor(0xF1, 0xF4, 0xF6) if r % 2 else RGBColor(0xFF, 0xFF, 0xFF))
    return t


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    light_blue, light_green, light_orange = (RGBColor(0xDB, 0xE8, 0xF4), RGBColor(0xDF, 0xF0, 0xE3),
                                             RGBColor(0xFB, 0xE3, 0xD2))

    # 1 Title and team (0:15)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    tf = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(3)).text_frame
    tf.word_wrap = True
    for i, (t, size, col) in enumerate([("AI-Powered Telecom Network Operations Assistant", 40, RGBColor(0xFF, 0xFF, 0xFF)),
                                        ("Similar-incident retrieval, multi-agent root-cause analysis and validated "
                                         "troubleshooting for NOC engineers", 20, RGBColor(0xDB, 0xE8, 0xF4)),
                                        ("(your name)  ·  GenAI capstone  ·  2026", 16, RGBColor(0xF6, 0xC9, 0xA8))]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text, r.font.size, r.font.name, r.font.color.rgb, r.font.bold = t, Pt(size), FONT, col, i == 0
        p.space_after = Pt(18)
    s.notes_slide.notes_text_frame.text = "0:15. Introduce yourself and the one-line pitch."

    # 2 Problem (0:45)
    s = slide(prs, "The problem: slow, inconsistent incident triage",
              "0:45. NOC engineers face the same degradations repeatedly; the knowledge is buried in KPI history.")
    text(s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12), Inches(5)).text_frame, [
        "NOC teams repeatedly face congestion, weak signal, latency, packet loss, call drops, handover failures and faults",
        "What similar incidents looked like, why they happened and what fixed them is buried in thousands of KPI records",
        "Searching by hand is slow → higher MTTR and inconsistent decisions between engineers",
        "Business need: from a plain-language description, return within seconds",
        "  similar past cases · a classified, prioritized incident · root-cause hypotheses · cited, validated actions",
        "And improve from engineer feedback"], size=22, space=12)

    # 3 Solution overview (0:30)
    s = slide(prs, "What the engineer gets", "0:30. One request returns all of this; every step is grounded and cited.")
    items = [("Similar cases", "hybrid search over 100,000 historical records, with filters", light_green, GREEN),
             ("Classification", "8 incident types + NORMAL; asks for clarification when vague", light_orange, ORANGE),
             ("Risk and priority", "ML fault risk, P1–P4 priority with reasons, impact level", light_green, GREEN),
             ("Root causes", "ranked hypotheses, each with evidence type and citations", light_orange, ORANGE),
             ("Actions", "ordered, phased, with verification and change-approval flags", light_orange, ORANGE),
             ("Validation", "an LLM judge + rule checks; one revision; low-confidence flag", light_orange, ORANGE)]
    for i, (t, b, f, l) in enumerate(items):
        box(s, 0.6 + (i % 3) * 4.15, 1.5 + (i // 3) * 2.55, 3.9, 2.25, t, [b], f, l, title_size=22, body_size=17)

    # 4 Architecture (0:45)
    s = slide(prs, "Architecture", "0:45. Top band: offline build, no LLM. Bottom band: the mandated spine "
              "Retriever → LLM/RAG Pipeline → Agents → API Layer → UI. Orange = LLM, green = deterministic.")
    pic_h = Inches(5.85)
    pic_w = int(pic_h * 2000 / 1255)
    s.shapes.add_picture(str(PROJECT_ROOT / "architecture" / "architecture.png"), int((W - pic_w) / 2), Inches(1.15),
                         height=pic_h)

    # 5 Dataset reality (0:40)
    s = slide(prs, "The dataset: three surprising findings", "0:40. These findings shaped the whole design.",
              subtitle="5G NetOps dataset (Hugging Face fenar/5g_netops): 100,000 rows × 27 KPI columns, Apache-2.0")
    box(s, 0.6, 1.8, 3.95, 4.9, "1. No text at all", [
        "No incident descriptions, causes or actions.", "", "→ Case narratives generated from the KPIs",
        "→ 9 authored runbooks as the knowledge source (81 sections)"], light_blue, NAVY, title_size=22, body_size=17)
    box(s, 4.7, 1.8, 3.95, 4.9, "2. Leaky columns", [
        "Fault occurrence rate is an exact formula of other columns (R² = 1.00).",
        "'Issue Reported' is noise (AUC 0.501).", "", "→ ML on leading indicators only:",
        "fault-risk AUC 0.874, honest not perfect"], light_blue, NAVY, title_size=22, body_size=17)
    box(s, 8.8, 1.8, 3.95, 4.9, "3. Synthetic", [
        "Faker city names (38,066 unique), uniform distributions, no timestamps or cell IDs.", "",
        "→ Severity relative to this network (p75 / p90)", "→ Trends cannot be verified: the app says so"],
        light_blue, NAVY, title_size=22, body_size=17)

    # 6 Tech stack (0:20)
    s = slide(prs, "Technology stack", "0:20. Everything runs locally except the LLM call; zero-key mode works too.")
    table(s, 0.6, 1.5, 12.1, [
        ["Layer", "Choice"],
        ["LLM", "Claude Haiku 5.5 via the AICredits gateway (≈ ₹0.75 per analysis); Sonnet 5.5 = best measured"],
        ["Agents", "LangGraph state graph: 4 agents, judge loop, deterministic fallbacks"],
        ["Retrieval", "ChromaDB (dense, bge-small-en-v1.5) + bm25s, reciprocal-rank fusion + rerank"],
        ["ML", "scikit-learn HistGradientBoosting: fault risk, throughput"],
        ["API", "FastAPI + Uvicorn: 13 endpoints, guardrails first, one error envelope"],
        ["Frontend", "React 19 + TypeScript (Vite, TanStack Query): live agent progress, summary-first report"],
        ["Storage", "SQLite: runs, feedback, LLM cache, spend ledger"],
        ["Evaluation", "DeepEval (LLM judge) + retrieval, classifier, guardrail, ML and judge tests; pytest"]],
        col_w=[2.2, 9.9], size=15, h_row=0.56)

    # 7 Agent workflow (0:35)
    s = slide(prs, "Agent workflow: why agents, and the judge loop",
              "0:35. Only LLM decisions under ambiguity are agents; everything computable is a deterministic tool. "
              "The judge's verdict drives one revision; a second failure is shown as a low-confidence diagnosis.")
    xs = [0.6, 3.75, 6.9, 10.05]
    for x, (t, b) in zip(xs, [("Classifier", ["category, symptoms,", "filters, scope"]),
                              ("Analyst", ["ranked root causes,", "evidence + citations"]),
                              ("Recommender", ["ordered actions,", "verification, approval"]),
                              ("Judge", ["4 scores + 6 rule checks", "pass / revise / fail"])]):
        box(s, x, 1.6, 2.7, 1.5, t, b, light_orange, ORANGE, title_size=18, body_size=14)
    for a, b in zip(xs, xs[1:]):
        arrow(s, a + 2.7, 2.35, b, 2.35)
    arrow(s, 11.4, 1.6, 5.1, 1.5, ORANGE, dashed=True)
    lbl = s.shapes.add_textbox(Inches(6.5), Inches(1.12), Inches(4), Inches(0.4)).text_frame
    r = lbl.paragraphs[0].add_run()
    r.text, r.font.size, r.font.name, r.font.color.rgb, r.font.italic = "revision (max 1)", Pt(13), FONT, ORANGE, True
    tools = ["KPI analyzer", "Segment statistics", "ML predictors", "Prioritization + impact", "Similar-case search",
             "Guardrails (before any LLM)"]
    for i, t in enumerate(tools):
        box(s, 0.6 + (i % 3) * 4.15, 3.7 + (i // 3) * 1.15, 3.9, 0.95, t, [], light_green, GREEN, title_size=17,
            centered=True)
    note = s.shapes.add_textbox(Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.8)).text_frame
    text(note, ["!Every LLM node has a deterministic fallback; with no key, or after the daily spend cap, the same request "
                "is answered by rules alone (reduced mode)."], size=15, color=GREY)

    # 8 Live demo (3:00)
    s = slide(prs, "Live demo (3 minutes)", "3:00. Turn on Demo mode (rail switch): larger type, five incident cards, "
              "keys 1 to 5. Repeated queries come from the response cache, so they are instant and free. "
              "Backup: the screenshot on the hidden slide, or zero-key mode.")
    table(s, 0.6, 1.5, 12.1, [
        ["Area", "What to show", "Time"],
        ["Investigate", "Preset 1: agents working live, then the report: diagnosis, priority, likely cause, "
         "escalation, worst KPIs, top 3 actions; open a tab (causes, similar cases, agent trace)", "1:40"],
        ["Cases", "Filters (state, season, weather) and the four-mode search comparison", "0:30"],
        ["Simulate", "Change KPIs → fault risk, throughput and priority update with reasons", "0:25"],
        ["Quality", "Judge verdicts, feedback loop, both evaluation runs named by model", "0:25"]],
        col_w=[2.2, 8.7, 1.2], size=15, h_row=0.8)
    tip = s.shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(12.1), Inches(1.0)).text_frame
    text(tip, ["A daily spend cap and the prepaid AICredits balance bound the total cost"], size=16, color=GREY)

    # 9 Evaluation (0:45)
    s = slide(prs, "Evaluation results",
              "0:45. Left: parts that use no model, identical in both runs. Right: the two configurations. Sonnet is "
              "the best measured capability; Haiku is what the demo runs, about 10x cheaper with the same accuracy, "
              "less consistent, and the judge flags it. Answer-quality scores come from different judges, so compare "
              "basic RAG vs multi-agent within a run, not across runs.",
              subtitle="Search, classification, guardrails, ML, answer quality vs basic RAG, judge reliability, cost")
    table(s, 0.6, 1.75, 5.2, [
        ["Same in both runs", "Result"],
        ["Search Hit@5 / nDCG@10", "1.000 / 0.953"],
        ["Guardrails blocked / false pos.", "12 of 12 / 0"],
        ["Fault-risk ROC-AUC", "0.874"],
        ["Priority scenarios", "6 of 6"],
        ["Feedback loop", "+0.006 nDCG, 0 worse"]], col_w=[3.2, 2.0], size=14, h_row=0.55)
    table(s, 6.1, 1.75, 6.6, [
        ["", "Sonnet 5.5", "Haiku 5.5 (demo)"],
        ["Targets met", "26 of 28", "27 of 28"],
        ["Classifier accuracy", "0.955", "0.955"],
        ["Actionability: RAG → agents", "0.60 → 0.84", "0.66 → 0.99"],
        ["Planted errors caught", "12 of 12", "12 of 12"],
        ["Median time per analysis", "31 s", "35 s"],
        ["Passed judge first time", "11 of 11", "8 of 11"],
        ["Cost per analysis", "≈ $0.076", "≈ $0.007"]], col_w=[2.8, 1.9, 1.9], size=14, h_row=0.55)

    # 10 Challenges and learnings (0:25)
    s = slide(prs, "Challenges and learnings", "0:25. End on the honest trade-offs.")
    text(s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12), Inches(5.3)).text_frame, [
        "Synthetic, text-free data → generated narratives, authored runbooks, leakage-free ML",
        "Agents cost 4 LLM calls instead of 1 → paid back in actionability and an independent check of every answer",
        "Provider migration (NVIDIA → Claude) → median latency 76 s → 31 s, no JSON failures",
        "Cost: Sonnet evaluation ≈ $11 → Haiku 5.5 via a prepaid gateway (≈ 10× cheaper per analysis),",
        "  same accuracy, less consistent: the judge revised 3 of 11 and flagged 2 low-confidence",
        "Information overload → a summary-first report: the essentials in 15 to 30 seconds, detail in tabs",
        "Next: real data with timestamps, retrieve runbook verification sections, authentication"],
        size=20, space=12)

    # Hidden backups
    s = slide(prs, "Backup: the report", "Hidden backup slide: use if the live demo fails. Recorded Haiku 5.5 run "
              "of PRD query 1 (32 s, 4 model calls, validated).")
    pic_h = Inches(5.9)
    pic_w = int(pic_h * 1440 / 900)
    s.shapes.add_picture(str(PROJECT_ROOT / "presentation" / "assets" / "report_top.png"), int((W - pic_w) / 2),
                         Inches(1.15), height=pic_h)
    s._element.set("show", "0")
    s = slide(prs, "Backup: Haiku 5.5 vs Sonnet 5.5", "Hidden backup slide. Full suite on both configurations.")
    table(s, 0.6, 1.5, 12.1, [
        ["Same test sets", "Haiku 5.5 via AICredits (demo)", "Sonnet 5.5 (best capability)"],
        ["Evaluation judge", "gpt-oss-120b", "Claude Opus 5.5"],
        ["Targets met", "27 of 28", "26 of 28"],
        ["Judge verdicts (11 queries)", "9 pass, 2 low confidence; 3 revised", "11 pass; 0 revised"],
        ["Median / max time per analysis", "34.8 s / 78.8 s", "31.4 s / 33.9 s"],
        ["Planted errors caught by the judge agent", "12 of 12, 0 false alarms", "12 of 12, 0 false alarms"],
        ["Cost of the full evaluation", "≈ $0.27", "≈ $11.4"]], col_w=[4.5, 3.8, 3.8], size=15, h_row=0.6)
    s._element.set("show", "0")
    s = slide(prs, "Backup: Q&A preparation", "Hidden backup slide.")
    text(s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12), Inches(5.3)).text_frame, [
        "Why does synthetic data matter? Generator artifacts; severity is relative; no trends",
        "How was leakage avoided? Leakage audit: R² from root columns ≥ 0.15 → excluded",
        "Why agents and not one call? Actionability 0.60 → 0.84 (Sonnet) and 0.66 → 0.99 (Haiku); judge catches 12/12",
        "Why Haiku for the demo? Same accuracy at ≈ 10x lower cost; weaker answers are revised or flagged, not hidden",
        "How is hallucination controlled? Citations validated against evidence IDs; rule checks; judge; fallbacks",
        "What does feedback change? A smoothed prior in the rerank (+0.006 nDCG on voted queries, 0 regressions)",
        "Why Chroma and BM25? Metadata filters on both paths; exact terms + paraphrases",
        "Production needs: real data, authentication, production vector DB, monitoring"], size=17, space=10)
    s._element.set("show", "0")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT.relative_to(PROJECT_ROOT)} ({len(prs.slides)} slides, 3 hidden)")


if __name__ == "__main__":
    main()
